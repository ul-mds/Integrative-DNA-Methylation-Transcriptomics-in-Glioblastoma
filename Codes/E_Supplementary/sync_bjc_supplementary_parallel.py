#!/usr/bin/env python3
"""
Synchronise BJC supplementary figures/tables with manuscript numbering.

Manuscript text cites Supplementary Figures 1–11 in sequential first-mention order.
BJC/Supplementary_Figures PNG files still use legacy numbering; this script renames
them so file N matches manuscript Supplementary Figure N.

Tables 1–23 already align across text, manuscript list, and BJC/Supplementary_Tables.
"""

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path("/Users/senol/Desktop/GBM/Methylation_Transcriptomics")
DOC = ROOT / "BJC Manuscript.docx"
BJC = ROOT / "BJC"
SF_DIR = BJC / "Supplementary_Figures"
ST_DIR = BJC / "Supplementary_Tables"
BACKUP = ROOT / "BJC Manuscript.backup_pre_bjc_sync.docx"

# Legacy file number -> manuscript (BJC submission) number
OLD_TO_NEW_SF = {1: 1, 2: 2, 3: 3, 4: 5, 5: 6, 6: 7, 7: 8, 8: 9, 9: 11, 10: 4, 11: 10}
NEW_TO_OLD_SF = {new: old for old, new in OLD_TO_NEW_SF.items()}

SF_LIST = {
    1: "Supplementary Figure 1. Survival analysis for cascade genes.",
    2: "Supplementary Figure 2. Detailed analysis of top cascade genes.",
    3: "Supplementary Figure 3. Spatial heterogeneity of DNA methylation in GB.",
    4: "Supplementary Figure 4. MGMT methylation profile similarity across the epigenome.",
    5: "Supplementary Figure 5. Age- and sex-specific DNA methylation dynamics in GB.",
    6: "Supplementary Figure 6. Methylation-based liquid biopsy candidate prioritisation for GB.",
    7: "Supplementary Figure 7. Methylation-immune crosstalk in GB.",
    8: "Supplementary Figure 8. Methylation-protein discordance patterns.",
    9: "Supplementary Figure 9. Internal and external validation of six exploratory prognostic "
       "biomarkers (B4GALT1, EGFR, VASP, GFPT2, PHF11, TSPYL5).",
    10: "Supplementary Figure 10. Integrative methylation–transcriptomics–proteomics study workflow for GB.",
    11: "Supplementary Figure 11. EGFR methylation in the context of amplification and mutation.",
}

ST_LIST = {
    1: "Supplementary Table 1: Dataset characteristics and summary statistics for methylation, mRNA, and protein data.",
    2: "Supplementary Table 2: Complete list of 1,678 genes exhibiting high-methylation → low-expression → low-protein cascade pattern with all correlation values and p-values.",
    3: "Supplementary Table 3: Pathway enrichment analysis for cascade genes showing enrichment in cell cycle regulation, apoptosis, and DNA repair pathways.",
    4: "Supplementary Table 4: Exploratory survival analysis for cascade genes with log-rank test p-values and group sizes (n = 97).",
    5: "Supplementary Table 5: Complete list of 83 biomarker candidates with high methylation and negative protein correlation.",
    6: "Supplementary Table 6: Detailed survival analysis for B4GALT1, EGFR, VASP, GFPT2, PHF11, and TSPYL5 including Cox regression results and multivariate analysis.",
    7: "Supplementary Table 7: Complete list of 842 age-dependent methylation genes in males with correlation values and p-values.",
    8: "Supplementary Table 8: Complete list of 1,456 age-dependent methylation genes in females with correlation values and p-values.",
    9: "Supplementary Table 9: Sex-specific differences analysis showing genes with significant sex-specific methylation patterns.",
    10: "Supplementary Table 10: Pathway enrichment analysis for age-sex-specific genes showing enrichment in aging-related pathways.",
    11: "Supplementary Table 11: Complete list of 50 liquid biopsy candidates with detectability scores, prognostic scores, and liquid biopsy scores.",
    12: "Supplementary Table 12: Clinical translation potential analysis showing detectability versus prognostic value for all candidates.",
    13: "Supplementary Table 13: Comparison with existing liquid biopsy approaches in GB showing advantages of methylation-based signatures.",
    14: "Supplementary Table 14: Complete list of discordant genes (232 genes) with methylation-protein correlation values and p-values.",
    15: "Supplementary Table 15: Pathway enrichment analysis for discordant genes showing enrichment in post-transcriptional regulation pathways.",
    16: "Supplementary Table 16: Post-transcriptional regulation mechanisms analysis including miRNA regulation, RNA-binding protein interactions, and protein stability control.",
    17: "Supplementary Table 17: Dataset characteristics and summary statistics including patient demographics, clinical information, and data quality metrics.",
    18: "Supplementary Table 18: TCGA internal validation results (70/30 train–test split) for six biomarkers.",
    19: "Supplementary Table 19: External validation results in GSE7696 and GSE230770.",
    20: "Supplementary Table 20: External validation summary by biomarker.",
    21: "Supplementary Table 21: EGFR genetic stratification analysis.",
    22: "Supplementary Table 22: Top 10 age-dependent methylation genes in males and females.",
    23: "Supplementary Table 23: Top 20 liquid biopsy candidates with detectability and prognostic scores.",
}


def hl(para, text):
    para.clear()
    run = para.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def extract_supp(text, kind):
    nums = []
    if kind == "fig":
        pat = re.compile(r"Supplementary\s+Figures?\s+(\d+)\s*(?:and|&|–|-|to)\s*(\d+)", re.I)
        single = re.compile(r"Supplementary\s+(?:Figure|Fig\.?)\s*(\d+)", re.I)
    else:
        pat = re.compile(r"Supplementary\s+Tables?\s+(\d+)\s*(?:and|&|–|-|to)\s*(\d+)", re.I)
        single = re.compile(r"Supplementary\s+Table\s*(\d+)", re.I)
    for m in pat.finditer(text):
        a, b = int(m.group(1)), int(m.group(2))
        nums.extend(range(min(a, b), max(a, b) + 1))
    for m in single.finditer(text):
        nums.append(int(m.group(1)))
    return nums


def audit(doc):
    ref_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "References")
    body_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Introduction")

    def first_order(kind):
        seen, order = set(), []
        for i in range(body_start, ref_start):
            t = doc.paragraphs[i].text
            if re.match(r"^(Table|Figure)\s+\d+\.", t.strip()):
                continue
            for n in extract_supp(t, kind):
                if n not in seen:
                    seen.add(n)
                    order.append(n)
        return order

    fo = first_order("fig")
    to = first_order("table")
    ok = lambda o: o == list(range(1, len(o) + 1))
    print("\n=== AUDIT ===")
    print(f"Supp Figures in text: {fo}  {'OK' if ok(fo) else 'FAIL'}")
    print(f"Supp Tables in text:  {to}  {'OK' if ok(to) else 'FAIL'}")


def renumber_sf_pngs(sf_dir: Path):
    # Backup legacy numbering before rename (for setup_bjc_folder rebuilds)
    legacy_backup = sf_dir / "_legacy_png_backup"
    if legacy_backup.exists():
        shutil.rmtree(legacy_backup)
    legacy_backup.mkdir()
    for f in sf_dir.glob("Supplementary_Figure_*.png"):
        shutil.copy2(f, legacy_backup / f.name)

    tmp = sf_dir / "_renumber_tmp"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir()

    for new, old in sorted(NEW_TO_OLD_SF.items()):
        src = sf_dir / f"Supplementary_Figure_{old}.png"
        if not src.exists():
            raise FileNotFoundError(f"Missing legacy figure: {src}")
        shutil.copy2(src, tmp / f"Supplementary_Figure_{new}.png")
        print(f"  SF{old} -> SF{new}: {src.name}")

    for f in sf_dir.glob("Supplementary_Figure_*.png"):
        f.unlink()

    for f in sorted(tmp.glob("*.png")):
        shutil.move(str(f), sf_dir / f.name)
    tmp.rmdir()
    print(f"  Renamed 11 PNG files in {sf_dir}")


def verify_table_files(st_dir: Path):
    missing = [n for n in range(1, 24) if not (st_dir / f"Supplementary_Table_{n}.xlsx").exists()]
    if missing:
        raise FileNotFoundError(f"Missing supplementary tables: {missing}")
    print(f"  Verified Supplementary_Table_1.xlsx … Supplementary_Table_23.xlsx in {st_dir}")


def update_manuscript_lists(doc):
    # Supplementary Tables list
    st_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Supplementary Tables")
    sf_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Supplementary Figures")

    for i in range(st_start + 1, sf_start):
        t = doc.paragraphs[i].text.strip()
        if t.startswith("Supplementary Table"):
            m = re.match(r"Supplementary Table (\d+)", t)
            if m:
                n = int(m.group(1))
                if n in ST_LIST and not t.startswith(ST_LIST[n].split("[")[0].strip()):
                    hl(doc.paragraphs[i], ST_LIST[n])
                    print(f"  Updated manuscript ST{n} list entry")

    # Supplementary Figures list
    for i in range(sf_start + 1, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t.startswith("Supplementary Figure"):
            m = re.match(r"Supplementary Figure (\d+)", t)
            if m:
                n = int(m.group(1))
                if n in SF_LIST:
                    hl(doc.paragraphs[i], SF_LIST[n])
        elif t in ("References", "Additional Information", ""):
            if t == "References":
                break


def update_readme():
    readme = BJC / "README.txt"
    text = readme.read_text(encoding="utf-8") if readme.exists() else ""
    sf_block = """Supplementary_Figures/
  Supplementary_Figure_1.png … 11.png   (numbered to match in-text citation order)
  SF1  Survival analysis for cascade genes
  SF2  Detailed analysis of top cascade genes
  SF3  Spatial heterogeneity of DNA methylation
  SF4  MGMT methylation profile similarity
  SF5  Age- and sex-specific methylation dynamics
  SF6  Liquid biopsy candidate prioritisation
  SF7  Methylation-immune crosstalk
  SF8  Methylation-protein discordance patterns
  SF9  Internal and external biomarker validation
  SF10 Integrative study workflow
  SF11 EGFR methylation (amplification/mutation context)
  Supplementary_Figures.docx            Combined supplementary figures document"""
    if "SF4  MGMT" not in text:
        import re as _re
        text = _re.sub(
            r"Supplementary_Figures/.*?Supplementary_Figures\.docx",
            sf_block,
            text,
            flags=_re.DOTALL,
        )
        readme.write_text(text, encoding="utf-8")
        print(f"  Updated {readme}")


def main():
    shutil.copy2(DOC, BACKUP)
    print(f"Backup: {BACKUP}")

    print("\n=== Renumber BJC Supplementary_Figures PNGs ===")
    renumber_sf_pngs(SF_DIR)

    print("\n=== Verify BJC Supplementary_Tables ===")
    verify_table_files(ST_DIR)

    print("\n=== Update manuscript ===")
    doc = Document(DOC)
    update_manuscript_lists(doc)
    doc.save(DOC)
    audit(Document(DOC))

    print("\n=== Copy manuscript to BJC/Manuscript ===")
    man_dir = BJC / "Manuscript"
    man_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(DOC, man_dir / "BJC Manuscript.docx")
    print(f"  Copied to {man_dir / 'BJC Manuscript.docx'}")

    update_readme()
    print(f"\nDone. Saved: {DOC}")


if __name__ == "__main__":
    main()
