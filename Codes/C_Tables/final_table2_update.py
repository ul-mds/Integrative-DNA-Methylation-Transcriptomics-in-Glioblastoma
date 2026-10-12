#!/usr/bin/env python3
"""
Update Table 2 with complete survival data
"""

from pathlib import Path

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[2]
doc_path = REPO_ROOT / "BJC" / "Revised BJC Manuscript.docx"
if not doc_path.exists():
    doc_path = REPO_ROOT / "BJC Manuscript.docx"

print("Loading Word document...")
doc = Document(doc_path)

# Complete Table 2 with survival data
table2_complete = """**Table 2. Top biomarker candidates with high methylation and negative protein correlation, including survival-significant genes.**

| Rank | Gene | Methylation Mean | Meth-Protein Correlation | Meth-Protein P-value | Survival Benefit (days) | Survival P-value | Biomarker Score |
|------|------|-----------------|------------------------|---------------------|------------------------|-----------------|----------------|
| 1 | CTSZ | 0.799 | -0.613 | 5.21e-11 | - | - | 2.50 |
| 2 | EGFR | 0.873 | -0.556 | 6.07e-9 | **+209** | **0.000551** | 2.26 |
| 3 | CCND2 | 0.729 | -0.545 | 1.37e-8 | - | - | 2.03 |
| 4 | B4GALT1 | 0.788 | -0.527 | 4.96e-8 | **+156** | **0.000245** | 1.96 |
| 5 | VASP | 0.715 | -0.465 | 2.39e-6 | **+283** | **0.001687** | 1.52 |
| 6 | PHF11 | 0.739 | -0.439 | 9.77e-6 | **+233** | **0.006932** | 1.41 |
| 7 | TSPYL5 | 0.722 | -0.432 | 1.34e-5 | **+224** | **0.008472** | 1.38 |
| 8 | GFPT2 | 0.826 | -0.395 | 8.12e-5 | **+184** | **0.002374** | 1.25 |"""

# Find Table 2
for i, para in enumerate(doc.paragraphs):
    if 'Table 2' in para.text and 'biomarker candidates' in para.text.lower():
        print(f"Found Table 2 header at paragraph {i}")
        
        # Check if table content follows (might be in same or next paragraph)
        # Look for table-like content in next few paragraphs
        table_start = None
        for j in range(i, min(i+5, len(doc.paragraphs))):
            para_check = doc.paragraphs[j]
            text_check = para_check.text.strip()
            if '|' in text_check and ('Rank' in text_check or 'Gene' in text_check):
                table_start = j
                break
        
        if table_start:
            # Replace table content
            para_table = doc.paragraphs[table_start]
            para_table.clear()
            para_table.add_run(table2_complete)
            print(f"Updated Table 2 at paragraph {table_start}")
        else:
            # Table might be in a Word table object, or we need to insert it
            # Insert after header
            new_para = doc.paragraphs[i].insert_paragraph_after()
            new_para.add_run(table2_complete)
            print(f"Inserted Table 2 after paragraph {i}")
        break

# Save document
doc.save(doc_path)
print(f"\n✓ Saved updated Word document: {doc_path}")
