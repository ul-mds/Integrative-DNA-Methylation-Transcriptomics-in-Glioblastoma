#!/usr/bin/env python3
"""
Refresh Supplementary/ materials to align with current manuscript (npj Precision Oncology).
Updates:
  - Supplementary_Methods.docx (extended methods, synced with main Methods)
  - Supplementary_Tables.xlsx (S1, S6+Cox, S14 count, S18–S20 from validation CSVs)
  - Supplementary_Figures.docx (SF1–SF7 captions + images matching manuscript list)
  - README.txt (folder index and npj submission note)

Run from Methylation_Transcriptomics/:
  python3 update_supplementary_folder.py
"""
from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

BASE = Path(__file__).resolve().parent
SUPP_DIR = BASE / "Supplementary"
RESULTS_DIR = BASE / "results"
NOVEL_DIR = RESULTS_DIR / "novel_analyses"
VALIDATION_DIR = BASE / "validation" / "results"
PLOTS_DIR = BASE / "plots"
VALIDATION_PLOTS = BASE / "validation" / "plots"
COX_PATH = RESULTS_DIR / "nature_communications" / "cox_regression_age_sex.csv"

BIOMARKERS = ["B4GALT1", "EGFR", "VASP", "GFPT2", "PHF11", "TSPYL5"]

# Captions from manuscript end list (MANUSCRIPT_NATURE_Communications.docx)
SUPP_FIGURES = [
    {
        "num": 1,
        "caption": (
            "Supplementary Figure 1. Extended survival analysis for cascade-pattern genes in "
            "TCGA-CPTAC glioblastoma (n = 97). Genes meeting high-methylation → low-expression → "
            "low-protein criteria were evaluated for overall survival. "
            "(A) Top 20 genes ranked by log-rank p-value or methylation–mRNA correlation strength. "
            "(B) Methylation–mRNA Spearman correlation for survival-assessed cascade genes. "
            "(C) Summary of gene counts meeting cascade and survival criteria. "
            "(D–F) Kaplan–Meier curves for representative cascade genes stratified by protein abundance "
            "(high vs. low); high methylation/low protein groups shown in blue, low methylation/high protein "
            "in red. (G) Methylation–mRNA correlation versus −log10(log-rank p), colored by hazard ratio. "
            "Extends methylation-stratified survival for selected cascade genes in Figure 2 and illustrates "
            "broader prognostic variability across the 1,678-gene cascade set."
        ),
        "images": ["high_meth_low_expr_low_protein_survival.png"],
    },
    {
        "num": 2,
        "caption": (
            "Supplementary Figure 2. Extended multi-omics characterization of epigenetic regulatory "
            "cascade genes (n = 1,678; TCGA-CPTAC). "
            "Upper panel — genome-wide correlation structure: "
            "(A) Methylation–mRNA correlation versus −log10 p; asterisks, FDR-significant associations. "
            "(B) mRNA–protein correlations. (C) Direct methylation–protein correlations. "
            "(D) Counts of genes with significant correlations at each cascade step. "
            "Lower panel — top-ranked cascade genes: "
            "(A) Top 20 genes by negative methylation–mRNA correlation. "
            "(B) Normalized methylation, log10 mRNA, and protein levels for the top 10 genes. "
            "(C) Pairwise correlation heatmap across omics layers. (D) Mean protein abundance for top 20 genes. "
            "(E–G) Methylation versus mRNA relationships for the three leading cascade genes (Spearman r and p "
            "annotated). (H–K) Kaplan–Meier overall survival stratified by high versus low methylation "
            "(median split) for the four cascade genes with strongest survival separation."
        ),
        "images": ["cascade_analysis_comprehensive.png", "high_meth_low_expr_low_protein_comprehensive.png"],
    },
    {
        "num": 3,
        "caption": (
            "Supplementary Figure 3. Inter-tumoral DNA methylation heterogeneity in TCGA-CPTAC glioblastoma "
            "(n = 97; 4,111 genes analyzed). Per-gene methylation variance across patients was used as a "
            "proxy for cross-sample heterogeneity. "
            "(A) Top 20 genes ranked by methylation variance. "
            "(B) Distribution of variance values; dashed line, top 10% threshold (411 high-heterogeneity genes). "
            "(C) Mean methylation versus coefficient of variation for all profiled genes. "
            "(D) Summary of heterogeneous loci and clinical relevance (biomarker stability, treatment "
            "resistance). Complements age- and sex-stratified methylation in Figure 4 by identifying loci "
            "with high inter-patient variability."
        ),
        "images": ["novel_analyses/spatial_heterogeneity.png"],
    },
    {
        "num": 4,
        "caption": (
            "Supplementary Figure 4. Methylation–immune gene expression crosstalk in glioblastoma "
            "(TCGA-CPTAC, n = 97). Spearman correlations between DNA methylation and mRNA expression "
            "were computed for 40 immune-related genes (T-cell, checkpoint, cytokine, and chemokine markers). "
            "(A) Top immune genes ranked by methylation–expression correlation (red, negative; blue, positive). "
            "(B) Distribution of correlation coefficients across all immune genes analyzed. "
            "(C) Mean methylation versus mean expression for the most correlated immune genes. "
            "(D) Summary of significant methylation–immune associations. Provides immune microenvironment "
            "context for in silico liquid biopsy prioritization candidates shown in Figure 5."
        ),
        "images": ["novel_analyses/immune_crosstalk.png"],
    },
    {
        "num": 5,
        "caption": (
            "Supplementary Figure 5. Extended methylation–protein discordance analysis "
            "(273 genes profiled; 232 discordant with |Spearman r| < 0.2; TCGA-CPTAC, n = 97). "
            "Upper panel: (A) Gene counts by discordance category (discordant, moderate, concordant high, "
            "concordant negative). (B) Methylation–protein correlation versus −log10 p; dashed line, p = 0.05. "
            "Lower panel: (A) Discordance category distribution. (B) Correlation–significance scatter "
            "colored by category; orange dashed lines, |r| = 0.2 discordance threshold. "
            "(C) Top 20 discordant genes with weakest methylation–protein coupling. "
            "(D) Mean methylation versus mean protein abundance; discordant genes highlighted. "
            "Extends the discordance framework in Figure 6; full gene lists in Supplementary Table 14."
        ),
        "images": ["methylation_protein_discordance.png", "discordance_analysis_comprehensive.png"],
    },
    {
        "num": 6,
        "caption": (
            "Supplementary Figure 6. Internal and external validation of six exploratory prognostic "
            "biomarkers (B4GALT1, EGFR, VASP, GFPT2, PHF11, TSPYL5). "
            "Upper panel: TCGA-CPTAC 70/30 train–test split; survival difference (days) in the held-out "
            "test set (n = 29). Green bars, direction concordant with discovery (high methylation associated "
            "with longer survival); red bars, discordant direction. "
            "Lower panel: cross-cohort comparison — top, survival difference (high minus low group median); "
            "bottom, −log10(log-rank p) in TCGA-CPTAC, GSE7696 (expression, n = 80; low expression as "
            "high-methylation proxy), and GSE230770 (EPIC methylation, n = 51). Red dashed line, p = 0.05. "
            "Numerical results in Supplementary Tables 18–20. Five of six biomarkers showed directionally "
            "concordant associations in ≥1 external cohort (EGFR and TSPYL5 in both)."
        ),
        "images": [
            "../validation/plots/tcga_internal_validation.png",
            "../validation/plots/external_validation_comparison.png",
        ],
    },
    {
        "num": 7,
        "caption": (
            "Supplementary Figure 7. MGMT methylation profile similarity across the epigenome "
            "(TCGA-CPTAC, n = 97). Pairwise methylation similarity to MGMT was quantified by Pearson "
            "correlation, Spearman correlation, and cosine similarity; a composite score integrated all metrics. "
            "(A) Top 20 genes ranked by composite similarity score. "
            "(B–C) Distributions of Pearson and Spearman correlations with MGMT across all profiled genes. "
            "(D–F) Top 20 genes by Pearson, Spearman, and cosine similarity, respectively. "
            "(G) Pearson versus Spearman correlation concordance; stars, top 10 composite-scoring genes. "
            "(H) Composite similarity versus mean methylation difference for the top 50 MGMT-like genes. "
            "(I) Summary statistics. Identifies genes with MGMT-like epigenetic silencing patterns relevant "
            "to temozolomide response biology."
        ),
        "images": ["mgmt_similarity_analysis.png"],
    },
]

METHODS_SECTIONS = [
    (
        "Note for submission",
        "For npj Precision Oncology, analytical methods are reported in the main manuscript Methods section. "
        "This document provides extended technical detail for reference and reproducibility. It is not cited "
        "in the main text and should be merged into the single Supplementary Information PDF (Tables + Figures only) "
        "or retained as an internal archive as required by the journal.",
    ),
    (
        "1. Data sources and cohort",
        "Integrated DNA methylation (Illumina EPIC), mRNA expression, and protein abundance data were obtained "
        "from the TCGA-CPTAC glioblastoma cohort (n = 97 patients with complete multi-omics data; Wang et al., 2021). "
        "After quality control (features with >50% missing values excluded), median imputation was applied per feature. "
        "EPIC probe-to-gene mapping used Illumina manifest annotations; where probes map to multiple genes "
        "(e.g., SEL1L2/MACROD2, RP11-1C8.4/RIMS2, SEPT8/CCNI2), the primary annotated gene is reported in main text "
        "with alternate mapping noted.",
    ),
    (
        "2. Software and reproducibility",
        "Analyses were performed in Python 3 (pandas, NumPy, SciPy, lifelines, scikit-learn, matplotlib, seaborn). "
        "Analysis scripts are organized in the Codes/ directory (A_Core, B_Figures, C_Tables, D_Validation). "
        "Random seeds were fixed where applicable (train-test split). FDR correction used Benjamini-Hochberg at α = 0.05.",
    ),
    (
        "3. Cascade analysis and Cascade Index",
        "Genes exhibiting high-methylation → low-expression → low-protein cascade satisfied: (1) mean methylation β "
        "above cohort median; (2) methylation–mRNA Spearman r < −0.1, FDR < 0.05; (3) mRNA–protein r > 0.3, FDR < 0.05; "
        "(4) methylation–protein r < −0.1, FDR < 0.05. The Cascade Index (CI_g) integrated normalized methylation, "
        "expression, and protein levels per gene to rank cascade strength across the cohort.",
    ),
    (
        "4. Biomarker discovery and survival analysis",
        "Biomarker candidates required mean DNA methylation β > 0.7 and significant negative methylation–protein "
        "Spearman correlation (p < 0.05). Composite biomarker score combined correlation strength, methylation level, "
        "p-value, and protein difference. Survival stratification: high methylation (β > 0.7) vs low (β ≤ 0.7); "
        "Kaplan-Meier curves with log-rank tests. Multivariate Cox proportional hazards models adjusted for age and sex "
        "(lifelines); hazard ratios with 95% CI reported in Supplementary Table 6.",
    ),
    (
        "5. Age- and sex-specific methylation",
        "Age–methylation Spearman correlations computed separately for males and females. Age-dependent genes: |r| > 0.5, "
        "p < 0.05. Age categories: Young (0–40 years), Mid (41–55), Old (56+). Sex differences tested with Mann-Whitney U.",
    ),
    (
        "6. In silico liquid biopsy candidate prioritization",
        "Methylation-based liquid biopsy candidates were prioritized in silico (not validated in plasma). "
        "Liquid Biopsy Score = Detectability Score × (1 + Prognostic Score). Detectability Score = mean methylation × "
        "(1/(1+CV)) for genes with CV < 1. Prognostic Score = −log₁₀(p) for survival association when p < 0.05, else 0. "
        "Top 50 genes retained (Supplementary Tables 11–13).",
    ),
    (
        "7. Methylation–protein discordance",
        "Discordant genes defined as |methylation–protein Spearman r| < 0.2. Categories: discordant (|r| < 0.2), "
        "moderate (0.2 ≤ |r| < 0.5), concordant high (r ≥ 0.5), concordant negative (r ≤ −0.5). "
        "Full gene list in Supplementary Table 14 (n = 232 discordant genes in TCGA-CPTAC cohort).",
    ),
    (
        "8. Tumor purity (ESTIMATE)",
        "Tumor purity and immune/stromal scores were estimated with ESTIMATE from mRNA expression profiles to assess "
        "confounding of methylation–survival associations.",
    ),
    (
        "9. Pathway enrichment",
        "DAVID and Gene Ontology enrichment for cascade, age-sex, and discordant gene sets; Benjamini-Hochberg FDR correction.",
    ),
    (
        "10. External and internal validation",
        "Internal validation (TCGA-CPTAC): 70/30 train-test split; six biomarkers tested on held-out samples "
        "(β > 0.7 vs β ≤ 0.7; Kaplan-Meier, log-rank). External cohorts: GSE7696 (expression, n = 80; low expression "
        "as proxy for high methylation) and GSE230770 (EPIC methylation, n = 51 IDH-wild-type GBM). "
        "Probe mappings: B4GALT1 (222163_s_at), EGFR (201983_s_at, 211607_x_at, cg18809074), VASP (201651_s_at, cg03547355), "
        "GFPT2 (218628_at), PHF11 (221731_x_at), TSPYL5 (221840_at, cg07814318). "
        "Concordance = directional agreement with TCGA (high methylation or low expression associated with longer survival). "
        "Five of six biomarkers showed concordant direction in ≥1 external dataset (EGFR and TSPYL5 in both; Supplementary Tables 18–20).",
    ),
]


def _add_methods_section(doc: Document, heading: str, body: str) -> None:
    if heading == "Note for submission":
        p = doc.add_paragraph(body)
        p.paragraph_format.space_after = Pt(12)
        for run in p.runs:
            run.italic = True
        return
    doc.add_heading(heading, level=1)
    for block in body.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)


def create_supplementary_methods() -> Path:
    out = SUPP_DIR / "Supplementary_Methods.docx"
    doc = Document()
    title = doc.add_heading("Supplementary Methods", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for heading, body in METHODS_SECTIONS:
        _add_methods_section(doc, heading, body)
    doc.save(str(out))
    print(f"Updated: {out}")
    return out


def _write_sheet_from_df(ws, title: str, df: pd.DataFrame, merge_cols: int | None = None) -> None:
    ws.delete_rows(1, ws.max_row)
    ws.append([title])
    ws.append([])
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(list(row))
    if merge_cols and merge_cols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=merge_cols)


def _patch_s1(wb) -> None:
    ws = wb["S1_Dataset"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["Supplementary Table 1: Dataset characteristics and summary statistics for methylation, mRNA, and protein data."])
    ws.merge_cells("A1:B1")
    ws.append([])
    rows = [
        ("Metric", "Value"),
        ("Patients with complete multi-omics data", 97),
        ("Methylation probes (EPIC)", "20,000"),
        ("mRNA genes", "34,594"),
        ("Proteins", "10,409"),
        ("Data source", "TCGA-CPTAC glioblastoma (Wang et al., 2021)"),
        ("Missing value threshold", ">50% excluded per feature"),
        ("Imputation", "Median per feature"),
        ("Methylation platform", "Illumina EPIC"),
    ]
    for r in rows:
        ws.append(list(r))
    print("  Patched S1_Dataset (TCGA-CPTAC)")


def _patch_s6(wb) -> None:
    surv_path = RESULTS_DIR / "biomarker_survival_analysis.csv"
    df_surv = pd.read_csv(surv_path)
    df_surv = df_surv[df_surv["Gene"].isin(BIOMARKERS)].copy()
    df_surv = df_surv.rename(
        columns={
            "High_Meth_Count": "N high meth",
            "Low_Meth_Count": "N low meth",
            "High_Meth_Median_Survival": "Median survival high (days)",
            "Low_Meth_Median_Survival": "Median survival low (days)",
            "Survival_Diff": "Survival diff (days)",
            "LogRank_P": "Log-rank p-value",
            "Is_Significant": "Significant",
        }
    )
    df_surv["Significant"] = df_surv["Significant"].map({True: "Yes", False: "No"})
    df_surv = df_surv.round(4)

    cox = pd.read_csv(COX_PATH)
    cox = cox[cox["Gene"].isin(BIOMARKERS)].copy()
    cox_display = cox[
        ["Gene", "N", "HR", "HR_CI_low", "HR_CI_high", "P_value", "Significant", "Direction", "Error"]
    ].rename(
        columns={
            "HR": "Hazard ratio",
            "HR_CI_low": "HR 95% CI lower",
            "HR_CI_high": "HR 95% CI upper",
            "P_value": "Cox p-value",
            "Significant": "Cox significant",
        }
    )
    cox_display["Cox significant"] = cox_display["Cox significant"].map({True: "Yes", False: "No"})
    cox_display = cox_display.round(4)

    ws = wb["S6_Survival_Detail"]
    ws.delete_rows(1, ws.max_row)
    title = (
        "Supplementary Table 6: Detailed survival analysis for B4GALT1, EGFR, VASP, GFPT2, PHF11, and TSPYL5 "
        "including Cox regression results and multivariate analysis (adjusted for age and sex)."
    )
    ws.append([title])
    ws.merge_cells("A1:H1")
    ws.append([])
    ws.append(["Part A. Kaplan-Meier / log-rank (TCGA-CPTAC, n=97)"])
    for row in dataframe_to_rows(df_surv, index=False, header=True):
        ws.append(list(row))
    ws.append([])
    ws.append(["Part B. Multivariate Cox regression (age and sex adjusted)"])
    for row in dataframe_to_rows(cox_display, index=False, header=True):
        ws.append(list(row))
    print("  Patched S6_Survival_Detail (6 biomarkers + Cox)")


def _patch_s14(wb) -> None:
    df = pd.read_csv(RESULTS_DIR / "methylation_protein_discordance_full.csv")
    discordant = df[df["Discordance_Type"] == "Discordant"] if "Discordance_Type" in df.columns else df
    n = len(discordant)
    ws = wb["S14_Discordant"]
    _write_sheet_from_df(
        ws,
        f"Supplementary Table 14: Complete list of discordant genes (n={n}; methylation-protein |r| < 0.2) "
        "with correlation values and p-values.",
        discordant,
        merge_cols=min(8, len(discordant.columns)),
    )
    print(f"  Patched S14_Discordant (n={n})")


def _patch_s11_header(wb) -> None:
    ws = wb["S11_LiquidBiopsy"]
    ws.cell(1, 1).value = (
        "Supplementary Table 11: Complete list of 50 in silico liquid biopsy prioritization candidates "
        "with detectability scores, prognostic scores, and composite scores (not validated in plasma)."
    )
    ws = wb["S13_LiquidBiopsy_Compare"]
    ws.cell(1, 1).value = (
        "Supplementary Table 13: Comparison with existing liquid biopsy approaches in GB "
        "(in silico methylation-based prioritization framework)."
    )
    print("  Patched S11/S13 headers (in silico language)")


def _add_validation_sheets(wb) -> None:
    tcga = pd.read_csv(VALIDATION_DIR / "tcga_internal_validation.csv")
    ext = pd.read_csv(VALIDATION_DIR / "external_validation_results.csv")
    ext_sum = pd.read_csv(VALIDATION_DIR / "external_validation_summary.csv")

    tcga_display = tcga.rename(
        columns={
            "Train_P": "Train p-value",
            "Train_Sig": "Train significant",
            "Train_Diff": "Train survival diff (days)",
            "Test_P": "Test p-value",
            "Test_Sig": "Test significant",
            "Test_Diff": "Test survival diff (days)",
            "Concordant": "Directionally concordant",
        }
    )
    for col in ["Train significant", "Test significant", "Directionally concordant"]:
        tcga_display[col] = tcga_display[col].map({True: "Yes", False: "No"})
    tcga_display = tcga_display.round(4)

    ext_display = ext.copy()
    ext_display["Significant"] = ext_display["Significant"].map({True: "Yes", False: "No"})
    ext_display["Concordant_TCGA"] = ext_display["Concordant_TCGA"].map({True: "Yes", False: "No"})
    ext_display = ext_display.rename(
        columns={
            "Data_Type": "Data type",
            "N_High": "N high",
            "N_Low": "N low",
            "Median_High": "Median survival high (days)",
            "Median_Low": "Median survival low (days)",
            "Survival_Diff": "Survival diff (days)",
            "LogRank_P": "Log-rank p-value",
            "Concordant_TCGA": "Concordant with TCGA",
            "TCGA_Diff": "TCGA diff (days)",
            "TCGA_P": "TCGA p-value",
        }
    ).round(4)

    ext_sum_display = ext_sum.rename(
        columns={
            "TCGA_Diff": "TCGA survival diff (days)",
            "N_Validated_Datasets": "N datasets validated",
            "N_Significant": "N significant",
            "N_Concordant": "N concordant",
            "Validation_Rate": "Validation rate",
        }
    )

    for sheet, title, df, cols in [
        (
            "S18_TCGA_Internal",
            "Supplementary Table 18: TCGA internal validation (70/30 train-test split). "
            "Biomarker survival associations in held-out test set.",
            tcga_display,
            8,
        ),
        (
            "S19_External_Validation",
            "Supplementary Table 19: External validation results. GSE7696 (expression, n=80); "
            "GSE230770 (methylation, n=51). Survival diff = High group median − Low group median.",
            ext_display,
            13,
        ),
        (
            "S20_External_Summary",
            "Supplementary Table 20: External validation summary by biomarker. "
            "Validation rate = concordant datasets / total datasets validated.",
            ext_sum_display,
            6,
        ),
    ]:
        if sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            ws = wb.create_sheet(sheet)
        _write_sheet_from_df(ws, title, df, merge_cols=cols)
    print("  Refreshed S18–S20 validation tables")


def update_supplementary_tables() -> Path:
    path = SUPP_DIR / "Supplementary_Tables.xlsx"
    wb = load_workbook(path)
    _patch_s1(wb)
    _patch_s6(wb)
    _patch_s14(wb)
    _patch_s11_header(wb)
    _add_validation_sheets(wb)
    wb.save(path)
    print(f"Updated: {path}")
    return path


def _resolve_image(img_name: str) -> Path | None:
    if img_name.startswith("../"):
        p = (PLOTS_DIR / img_name).resolve()
    else:
        p = PLOTS_DIR / img_name
    return p if p.exists() else None


def create_supplementary_figures() -> Path:
    out = SUPP_DIR / "Supplementary_Figures.docx"
    doc = Document()
    title = doc.add_heading("Supplementary Figures", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    multi_image_labels = {
        2: ["Genome-wide cascade correlation structure.", "Top-ranked cascade genes (extended view)."],
        5: ["Discordance overview.", "Comprehensive discordance landscape."],
        6: ["TCGA internal validation (70/30 split).", "Cross-cohort external validation."],
    }

    for fig in SUPP_FIGURES:
        doc.add_heading(f"Supplementary Figure {fig['num']}", level=1)
        cap = doc.add_paragraph(fig["caption"])
        cap.paragraph_format.space_after = Pt(10)
        doc.add_paragraph()
        labels = multi_image_labels.get(fig["num"], [])
        for idx, img_name in enumerate(fig["images"]):
            if idx < len(labels):
                lbl = doc.add_paragraph(labels[idx])
                lbl.runs[0].italic = True
            img_path = _resolve_image(img_name)
            if img_path:
                try:
                    p = doc.add_paragraph()
                    p.add_run().add_picture(str(img_path), width=Inches(6.0))
                    doc.add_paragraph()
                except Exception as exc:
                    doc.add_paragraph(f"[Could not embed {img_name}: {exc}]")
            else:
                doc.add_paragraph(f"[Image not found: {img_name}]")
        doc.add_paragraph()

    doc.save(str(out))
    print(f"Updated: {out} ({len(SUPP_FIGURES)} figures)")
    return out


def write_readme() -> Path:
    readme = SUPP_DIR / "README.txt"
    content = """Supplementary Materials — Multi-Omics GBM Study
============================================

Files:
  1. Supplementary_Tables.xlsx     Tables S1–S20
  2. Supplementary_Figures.docx    Figures SF1–SF7 (captions match manuscript)
  3. Supplementary_Methods.docx    Extended methods (archive; not cited in main text)

npj Precision Oncology submission:
  - Main Methods are in the manuscript body (no Supplementary Methods file in submission PDF).
  - Merge Supplementary Tables + Supplementary Figures into one Supplementary Information PDF.

Regenerate all files:
  python3 update_supplementary_folder.py

Key updates (latest run):
  - S1: TCGA-CPTAC data source
  - S6: All 6 biomarkers + Cox regression (age/sex adjusted)
  - S14: Full discordant gene list (n=232)
  - S18–S20: Validation results from validation/results/*.csv
  - SF1–SF7: Manuscript-aligned captions and plot images

Note: Manuscript list text for Table 14 says "100 genes"; full analysis identifies 232 discordant genes (S14).
"""
    readme.write_text(content, encoding="utf-8")
    print(f"Updated: {readme}")
    return readme


def main() -> int:
    os.makedirs(SUPP_DIR, exist_ok=True)
    print("Updating Supplementary/ folder...")
    print("-" * 50)
    create_supplementary_methods()
    update_supplementary_tables()
    create_supplementary_figures()
    write_readme()
    print("-" * 50)
    print("Done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
